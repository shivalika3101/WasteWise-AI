from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "deliverables" / "ui-prototype"
OUTPUT_DOCX = OUTPUT_DIR / "WasteWise_AI_UI_Prototype.docx"
OUTPUT_PDF = OUTPUT_DIR / "WasteWise_AI_UI_Prototype.pdf"
SCREENSHOT_DIR = OUTPUT_DIR / "screenshots"


GREEN_DARK = RGBColor(0x1B, 0x43, 0x32)
GREEN_MID = RGBColor(0x2E, 0x7D, 0x32)
GREEN_SOFT = RGBColor(0xD8, 0xF3, 0xDC)
TEXT_MUTED = RGBColor(0x66, 0x66, 0x66)
GRAY_LIGHT = "D9E2DC"


SCREENS = [
    {
        "name": "Home",
        "route": "/",
        "purpose": "Primary landing page for the prototype. It introduces the WasteWise AI concept and keeps the first step focused on image upload or live camera capture.",
        "status": "Implemented and captured",
        "sections": [
            "Header with Cardiff University Project branding and two navigation states: Home and Bin Classification.",
            "Large editorial hero message with two calls to action: Browse Files and Live Scan.",
            "Stat cards highlighting UK categories, simplified home-page actions, and the two-page prototype structure.",
            "Upload and capture section with a dedicated action box for Browse Files and Live Scan.",
            "Planned system features grid covering image upload, live camera capture, bin classification, and UK guidance.",
        ],
        "actions": [
            "Browse Files",
            "Live Scan",
            "Switch to Bin Classification view",
        ],
        "notes": [
            "The running website uses a more polished design than the earlier source files that were open in the editor.",
            "This capture reflects the built prototype currently served from the local preview environment.",
        ],
        "screenshot": SCREENSHOT_DIR / "home-full.png",
        "review_focus": "Branding clarity, first-step guidance, hero readability, and whether upload vs live scan is obvious.",
    },
    {
        "name": "Bin Classification",
        "route": "Client-side bin state",
        "purpose": "Dedicated second view for bin-decision behaviour. It simulates whether the selected UK bin should open or remain closed based on the scanned item.",
        "status": "Implemented and captured",
        "sections": [
            "Intro block with a clear explanation of the separate bin-classification workflow.",
            "Live bin scanner panel with camera frame, status text, and capture controls.",
            "Decision output panel with selectable UK waste-bin categories.",
            "Animated bin cards previewing lid behaviour by category.",
            "Awaiting-scan result card ready to show the next decision state.",
        ],
        "actions": [
            "Open Live Scanner",
            "Capture Image",
            "Select Current Bin",
            "Analyse Bin Choice",
        ],
        "notes": [
            "This screen is toggled client-side from the Home navigation rather than via a separate URL.",
            "The captured state shows the live-scanner mode active with camera access still pending in headless review.",
        ],
        "screenshot": SCREENSHOT_DIR / "bin-classification.png",
        "review_focus": "Decision flow clarity, bin selection usability, and whether the second page feels distinct enough from the home page.",
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_text(run, size, bold=False, color=None):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_paragraph(doc, text="", style=None, align=None):
    p = doc.add_paragraph(style=style)
    if align:
        p.alignment = align
    if text:
        run = p.add_run(text)
        style_text(run, 12)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(item)
        style_text(run, 11, color=TEXT_MUTED)


def add_screen_table(doc, screen):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(1.7)
    table.columns[1].width = Inches(4.8)

    headers = table.rows[0].cells
    headers[0].text = "Field"
    headers[1].text = "Details"
    for cell in headers:
        set_cell_shading(cell, "D8F3DC")
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                style_text(run, 11, bold=True, color=GREEN_DARK)

    rows = [
        ("Route", screen["route"]),
        ("Purpose", screen["purpose"]),
        ("Status", screen["status"]),
        ("Primary actions", ", ".join(screen["actions"])),
        ("Key sections", "\n".join(screen["sections"])),
        ("Notes", "\n".join(screen["notes"])),
    ]

    for label, value in rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value
        for index, cell in enumerate(row):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if index == 0:
                set_cell_shading(cell, "F4FBF6")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    style_text(run, 10.5, bold=index == 0, color=GREEN_DARK if index == 0 else None)


def add_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9E2DC")
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def add_docx_screenshot(doc, image_path, width=6.5):
    if not image_path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    p.paragraph_format.space_after = Pt(8)

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(f"Captured from local preview: {image_path.name}")
    style_text(run, 9, color=GREEN_MID)


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(13)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("WasteWise AI UI Prototype")
    style_text(run, 22, bold=True, color=GREEN_DARK)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Current frontend review pack for Team B")
    style_text(run, 12, color=GREEN_MID)

    meta = doc.add_table(rows=3, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Table Grid"
    meta.autofit = False
    meta.columns[0].width = Inches(2.1)
    meta.columns[1].width = Inches(4.7)
    meta_rows = [
        ("Prepared", "16 July 2026"),
        ("Source", "Local Vite preview of the current built prototype in frontend/dist"),
        ("Purpose", "Share the current UI flow, actual rendered screens, and review notes with team members."),
    ]
    for i, (label, value) in enumerate(meta_rows):
        row = meta.rows[i].cells
        row[0].text = label
        row[1].text = value
        set_cell_shading(row[0], "D8F3DC")
        for cell in row:
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    style_text(run, 10.5, bold=cell is row[0], color=GREEN_DARK if cell is row[0] else None)

    add_paragraph(doc)
    overview = doc.add_paragraph()
    overview.paragraph_format.space_after = Pt(10)
    text = (
        "This document captures the current WasteWise AI frontend prototype using real screenshots from the "
        "running local preview. It is intended as a quick review artifact for team discussion, UI feedback, "
        "and next-step planning."
    )
    run = overview.add_run(text)
    style_text(run, 11, color=TEXT_MUTED)

    heading = doc.add_paragraph(style="Heading 1")
    run = heading.add_run("Prototype Flow")
    style_text(run, 16, bold=True, color=GREEN_DARK)

    flow = doc.add_paragraph()
    flow.paragraph_format.space_after = Pt(12)
    run = flow.add_run(
        "Home landing page -> Browse Files or Live Scan entry point\n"
        "Home navigation -> Bin Classification page\n"
        "Bin Classification page -> Capture item -> Analyse bin choice"
    )
    style_text(run, 11, color=TEXT_MUTED)

    heading = doc.add_paragraph(style="Heading 1")
    run = heading.add_run("Captured Screens")
    style_text(run, 16, bold=True, color=GREEN_DARK)
    add_paragraph(
        doc,
        "The screenshots below were captured from the live local preview so the document matches the current website visual state.",
    )
    add_docx_screenshot(doc, SCREENSHOT_DIR / "home.png", width=6.2)
    add_docx_screenshot(doc, SCREENSHOT_DIR / "bin-classification.png", width=6.2)

    heading = doc.add_paragraph(style="Heading 1")
    run = heading.add_run("Visual Direction")
    style_text(run, 16, bold=True, color=GREEN_DARK)
    add_bullets(
        doc,
        [
            "Primary style direction is eco-focused: green palette, white surfaces, soft gray shadows, and rounded cards.",
            "Typography uses Poppins in the frontend, with large hero headings and medium-weight section titles.",
            "Interaction style is friendly and simple, with prominent calls to action and hover lift effects.",
            "Core recurring patterns are centered cards, soft backgrounds, strong section spacing, and bold action buttons.",
        ],
    )

    heading = doc.add_paragraph(style="Heading 1")
    run = heading.add_run("Screen Summary")
    style_text(run, 16, bold=True, color=GREEN_DARK)

    summary = doc.add_table(rows=1, cols=4)
    summary.alignment = WD_TABLE_ALIGNMENT.LEFT
    summary.style = "Table Grid"
    summary.autofit = False
    widths = [1.8, 1.2, 1.5, 2.7]
    for index, width in enumerate(widths):
        summary.columns[index].width = Inches(width)

    header_cells = summary.rows[0].cells
    for cell, label in zip(header_cells, ["Screen", "Route", "Status", "Review focus"]):
        cell.text = label
        set_cell_shading(cell, "D8F3DC")
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                style_text(run, 10.5, bold=True, color=GREEN_DARK)

    for screen in SCREENS:
        row = summary.add_row().cells
        row[0].text = screen["name"]
        row[1].text = screen["route"]
        row[2].text = screen["status"]
        row[3].text = screen["review_focus"]
        for index, cell in enumerate(row):
            set_cell_margins(cell)
            if index == 0:
                set_cell_shading(cell, "F4FBF6")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    style_text(run, 10, bold=index == 0, color=GREEN_DARK if index == 0 else None)

    for screen in SCREENS:
        doc.add_section(WD_SECTION.NEW_PAGE)
        heading = doc.add_paragraph(style="Heading 1")
        run = heading.add_run(screen["name"])
        style_text(run, 16, bold=True, color=GREEN_DARK)

        sub = doc.add_paragraph()
        sub.paragraph_format.space_after = Pt(8)
        run = sub.add_run(f"Prototype screen reference: {screen['route']}")
        style_text(run, 10.5, color=GREEN_MID)

        add_docx_screenshot(doc, screen["screenshot"])
        add_screen_table(doc, screen)

        h2 = doc.add_paragraph(style="Heading 2")
        run = h2.add_run("Team Review Prompts")
        style_text(run, 13, bold=True, color=GREEN_DARK)

        prompts = {
            "Home": [
                "Is the home page clear enough about the difference between Browse Files and Live Scan?",
                "Do the stats and feature cards help explain the prototype, or do they add too much noise before action?",
                "Is the typography strong and readable enough for a first impression on desktop and mobile?",
            ],
            "Bin Classification": [
                "Does moving bin logic to a dedicated page make the workflow easier to understand?",
                "Are the UK bin labels and visual colors clear enough without extra explanation?",
                "Should the bin decision result be even more prominent once a scan has been analysed?",
            ],
        }
        add_bullets(doc, prompts[screen["name"]])
        add_divider(doc)

    doc.add_section(WD_SECTION.NEW_PAGE)
    heading = doc.add_paragraph(style="Heading 1")
    run = heading.add_run("Current Gaps and Notes")
    style_text(run, 16, bold=True, color=GREEN_DARK)
    add_bullets(
        doc,
        [
            "The built prototype currently behaves as a two-state single-page experience rather than the older multi-route flow found in some source files.",
            "Camera access is required for the live scan and live bin scanner paths; in headless review those states remain in a request/pending condition.",
            "Bin Classification is toggled internally from the navigation instead of loading on a dedicated URL path.",
            "The repo README still notes that backend integration is planned for a future stage.",
            "An open IDE tab referenced Contact.jsx, but that file does not exist in the current project tree.",
            "An open IDE tab referenced Loader.jsx, but the loading UI is currently implemented in Loading.jsx instead.",
        ],
    )

    heading = doc.add_paragraph(style="Heading 1")
    run = heading.add_run("Recommended Next Steps")
    style_text(run, 16, bold=True, color=GREEN_DARK)
    add_bullets(
        doc,
        [
            "Collect team feedback on whether the current two-page structure is easier to understand than the earlier multi-screen flow.",
            "Decide whether Bin Classification should remain a client-side view switch or become a true route with a shareable URL.",
            "Add explicit success, failure, and permission-denied states for camera-based actions.",
            "Connect the prototype to real classification output so the decision and result cards reflect live data.",
        ],
    )

    doc.save(OUTPUT_DOCX)


def build_pdf():
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def add_pdf_screenshot(story, image_path, max_width=6.1 * inch, max_height=7.2 * inch):
        if not image_path.exists():
            return
        reader = ImageReader(str(image_path))
        img_width, img_height = reader.getSize()
        scale = min(max_width / img_width, max_height / img_height)
        story.append(Image(str(image_path), width=img_width * scale, height=img_height * scale))
        story.append(Spacer(1, 0.08 * inch))
        story.append(Paragraph(f"Captured from local preview: {image_path.name}", styles["WWCaption"]))
        story.append(Spacer(1, 0.12 * inch))

    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="WWTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=20,
            textColor=colors.HexColor("#1B4332"),
            alignment=TA_CENTER,
            spaceAfter=10,
        )
    )
    styles.add(
        ParagraphStyle(
            name="WWSubtitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=11,
            textColor=colors.HexColor("#2E7D32"),
            alignment=TA_CENTER,
            spaceAfter=14,
        )
    )
    styles.add(
        ParagraphStyle(
            name="WWHeading",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=14,
            textColor=colors.HexColor("#1B4332"),
            spaceAfter=8,
        )
    )
    styles.add(
        ParagraphStyle(
            name="WWBody",
            parent=styles["BodyText"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#4F4F4F"),
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="WWCaption",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8.5,
            textColor=colors.HexColor("#2E7D32"),
            alignment=TA_CENTER,
            spaceAfter=8,
        )
    )

    doc = SimpleDocTemplate(
        str(OUTPUT_PDF),
        pagesize=letter,
        topMargin=0.7 * inch,
        bottomMargin=0.6 * inch,
        leftMargin=0.8 * inch,
        rightMargin=0.8 * inch,
    )
    story = []

    story.append(Paragraph("WasteWise AI UI Prototype", styles["WWTitle"]))
    story.append(Paragraph("Current frontend review pack for Team B", styles["WWSubtitle"]))

    meta_data = [
        ["Prepared", "16 July 2026"],
        ["Source", "Local Vite preview of the current built prototype in frontend/dist"],
        ["Purpose", "Share the current UI flow, actual rendered screens, and review notes with team members."],
    ]
    meta = Table(meta_data, colWidths=[1.4 * inch, 4.9 * inch])
    meta.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#D8F3DC")),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#1B4332")),
                ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                ("FONTNAME", (1, 0), (1, -1), "Helvetica"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9E2DC")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("PADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.extend([meta, Spacer(1, 0.2 * inch)])

    story.append(
        Paragraph(
            "This document captures the current WasteWise AI frontend prototype using real screenshots from the running local preview. "
            "It is intended as a quick review artifact for team discussion, UI feedback, and next-step planning.",
            styles["WWBody"],
        )
    )
    story.append(Paragraph("Prototype Flow", styles["WWHeading"]))
    story.append(
        Paragraph(
            "Home landing page -> Browse Files or Live Scan entry point<br/>"
            "Home navigation -> Bin Classification page<br/>"
            "Bin Classification page -> Capture item -> Analyse bin choice",
            styles["WWBody"],
        )
    )
    story.append(Paragraph("Captured Screens", styles["WWHeading"]))
    story.append(
        Paragraph(
            "The screenshots below were captured from the live local preview so the document matches the current website visual state.",
            styles["WWBody"],
        )
    )
    add_pdf_screenshot(story, SCREENSHOT_DIR / "home.png", max_height=4.2 * inch)
    add_pdf_screenshot(story, SCREENSHOT_DIR / "bin-classification.png", max_height=4.8 * inch)
    story.append(Paragraph("Visual Direction", styles["WWHeading"]))
    for bullet in [
        "Primary style direction is eco-focused: green palette, white surfaces, soft gray shadows, and rounded cards.",
        "Typography uses Poppins in the frontend, with large hero headings and medium-weight section titles.",
        "Interaction style is friendly and simple, with prominent calls to action and hover lift effects.",
        "Core recurring patterns are centered cards, soft backgrounds, strong section spacing, and bold action buttons.",
    ]:
        story.append(Paragraph(f"- {bullet}", styles["WWBody"]))

    story.append(Paragraph("Screen Summary", styles["WWHeading"]))
    summary_rows = [["Screen", "Route", "Status", "Review focus"]]
    for screen in SCREENS:
        summary_rows.append([screen["name"], screen["route"], screen["status"], screen["review_focus"]])

    summary = Table(summary_rows, colWidths=[1.2 * inch, 1.0 * inch, 1.7 * inch, 2.5 * inch])
    summary.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D8F3DC")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1B4332")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9E2DC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("PADDING", (0, 0), (-1, -1), 5),
                ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#F4FBF6")),
            ]
        )
    )
    story.extend([summary, PageBreak()])

    for index, screen in enumerate(SCREENS):
        story.append(Paragraph(screen["name"], styles["WWHeading"]))
        story.append(Paragraph(f"Prototype screen reference: {screen['route']}", styles["WWBody"]))
        add_pdf_screenshot(story, screen["screenshot"])
        rows = [
            ["Field", "Details"],
            ["Route", screen["route"]],
            ["Purpose", screen["purpose"]],
            ["Status", screen["status"]],
            ["Primary actions", ", ".join(screen["actions"])],
            ["Key sections", "<br/>".join(screen["sections"])],
            ["Notes", "<br/>".join(screen["notes"])],
        ]
        table = Table(rows, colWidths=[1.5 * inch, 4.7 * inch])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D8F3DC")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1B4332")),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
                    ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#F4FBF6")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9E2DC")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("PADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        story.extend([table, Spacer(1, 0.15 * inch), Paragraph("Team Review Prompts", styles["WWHeading"])])
        prompts = {
            "Home": [
                "Is the home page clear enough about the difference between Browse Files and Live Scan?",
                "Do the stats and feature cards help explain the prototype, or do they add too much noise before action?",
                "Is the typography strong and readable enough for a first impression on desktop and mobile?",
            ],
            "Bin Classification": [
                "Does moving bin logic to a dedicated page make the workflow easier to understand?",
                "Are the UK bin labels and visual colors clear enough without extra explanation?",
                "Should the bin decision result be even more prominent once a scan has been analysed?",
            ],
        }
        for bullet in prompts[screen["name"]]:
            story.append(Paragraph(f"- {bullet}", styles["WWBody"]))
        if index != len(SCREENS) - 1:
            story.append(PageBreak())

    story.append(PageBreak())
    story.append(Paragraph("Current Gaps and Notes", styles["WWHeading"]))
    for bullet in [
        "The built prototype currently behaves as a two-state single-page experience rather than the older multi-route flow found in some source files.",
        "Camera access is required for the live scan and live bin scanner paths; in headless review those states remain in a request/pending condition.",
        "Bin Classification is toggled internally from the navigation instead of loading on a dedicated URL path.",
        "The repo README still notes that backend integration is planned for a future stage.",
        "An open IDE tab referenced Contact.jsx, but that file does not exist in the current project tree.",
        "An open IDE tab referenced Loader.jsx, but the loading UI is currently implemented in Loading.jsx instead.",
    ]:
        story.append(Paragraph(f"- {bullet}", styles["WWBody"]))

    story.append(Paragraph("Recommended Next Steps", styles["WWHeading"]))
    for bullet in [
        "Collect team feedback on whether the current two-page structure is easier to understand than the earlier multi-screen flow.",
        "Decide whether Bin Classification should remain a client-side view switch or become a true route with a shareable URL.",
        "Add explicit success, failure, and permission-denied states for camera-based actions.",
        "Connect the prototype to real classification output so the decision and result cards reflect live data.",
    ]:
        story.append(Paragraph(f"- {bullet}", styles["WWBody"]))

    doc.build(story)


if __name__ == "__main__":
    build_document()
    try:
        build_pdf()
        print(f"Created {OUTPUT_PDF}")
    except ModuleNotFoundError:
        print("Skipped PDF generation because reportlab is not installed.")
    print(f"Created {OUTPUT_DOCX}")
